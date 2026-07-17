import io
from django.core.files.base import ContentFile
from openpyxl import Workbook
from regulatory.registry import agency_registry
from base.models import RegulatoryReportDocument


def render_xlsx(generated) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for section in generated.sections:
        # Excel sheet name limit
        ws = wb.create_sheet(title=section.title[:31])
        ws.append(section.headers)
        for row in section.rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_and_attach(report, uploaded_by) -> RegulatoryReportDocument:
    backend = agency_registry.get(report.agency)
    if not backend:
        raise ValueError(
            f"No backend registered for agency '{report.agency}'."
        )

    generated = backend.generate_report(report)

    if generated is None:
        raise ValueError(
            f"{backend.agency_name} has no generator for report type '{report.report_type}'."
        )

    content = render_xlsx(generated)

    doc = RegulatoryReportDocument(
        report=report,
        description=f"Auto-generated {generated.filename}",
    )

    doc.file.save(
        generated.filename,
        ContentFile(content),
        save=False
    )
    doc.full_clean()   # runs the magic-number check from earlier
    doc.save()
    return doc
